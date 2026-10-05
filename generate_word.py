import re
import json

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx2pdf import convert

FONT = 'Times New Roman'
ML_PREFIX = "Data Analysis & Machine Learning Projects - "

# סדר האלמנטים בסכמה של Word (כדי להכניס תגיות במקום הנכון)
PPR_AFTER_BIDI = (
    'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing',
    'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
    'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
    'w:pPrChange')
RPR_AFTER_BCS = (
    'w:i', 'w:iCs', 'w:caps', 'w:smallCaps', 'w:strike', 'w:dstrike', 'w:outline',
    'w:shadow', 'w:emboss', 'w:imprint', 'w:noProof', 'w:snapToGrid', 'w:vanish',
    'w:webHidden', 'w:color', 'w:spacing', 'w:w', 'w:kern', 'w:position', 'w:sz',
    'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd', 'w:fitText',
    'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout',
    'w:specVanish', 'w:oMath')
RPR_AFTER_SZCS = RPR_AFTER_BCS[RPR_AFTER_BCS.index('w:highlight'):]
RPR_AFTER_RTL = RPR_AFTER_BCS[RPR_AFTER_BCS.index('w:cs'):]


def clean(text):
    """מסיר נקודה בסוף שורה (נקודות באמצע משפט נשארות)"""
    text = text.strip()
    if text.endswith('.'):
        text = text[:-1].rstrip()
    return text


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = docx.oxml.shared.OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = docx.oxml.shared.OxmlElement('w:r')
    rPr = docx.oxml.shared.OxmlElement('w:rPr')

    rFonts = docx.oxml.shared.OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    rPr.append(rFonts)

    c = docx.oxml.shared.OxmlElement('w:color')
    c.set(qn('w:val'), '0000FF')
    rPr.append(c)
    u = docx.oxml.shared.OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)

    new_run.append(rPr)
    text_el = docx.oxml.shared.OxmlElement('w:t')
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_tnr_run(paragraph, text, bold=False, size=10):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:cs'), FONT)
    return run


def add_tnr_heading(doc, text, level=1, space_before=4):
    h = doc.add_heading(level=level)
    h.text = ""
    run = add_tnr_run(h, text, bold=True, size=11.5)
    run.font.color.rgb = RGBColor(46, 116, 181)
    h.paragraph_format.space_before = Pt(space_before)
    h.paragraph_format.space_after = Pt(2)
    return h


def add_bullet_point(doc, text, is_hebrew):
    text = clean(text)
    if is_hebrew:
        # הזחה תלויה: הבולט בולט החוצה והטקסט המשובר מתיישר מתחת לטקסט
        bp = doc.add_paragraph()
        bp.paragraph_format.left_indent = Inches(0.25)          # ב-bidi זה הצד הימני
        bp.paragraph_format.first_line_indent = Inches(-0.18)
        add_tnr_run(bp, "•\t" + text)
    else:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.left_indent = Inches(0.25)
        add_tnr_run(bp, text)


def add_numbered_item(doc, name, desc):
    name = name.rstrip(':')
    m = re.match(r'^(\d+\.)\s*(.*)$', name)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    if m:
        p.paragraph_format.first_line_indent = Inches(-0.2)
        add_tnr_run(p, m.group(1) + "\t")
        name = m.group(2)
    add_tnr_run(p, f"{name}: ", bold=True)
    add_tnr_run(p, desc)   # כאן הנקודה בסוף נשארת, כמו בקובץ ההדוגמה


def _insert(parent, tag, successors, val=None):
    el = docx.oxml.shared.OxmlElement(tag)
    if val is not None:
        el.set(qn('w:val'), str(val))
    parent.insert_element_before(el, *successors)


def make_rtl(p):
    """כיווניות ימין-לשמאל לפסקה: bidi, בלי יישור מפורש, ו-rtl/bCs/szCs לכל ריצה"""
    pPr = p._element.get_or_add_pPr()
    _insert(pPr, 'w:bidi', PPR_AFTER_BIDI, 1)
    p.alignment = None   # ב-bidi ברירת המחדל היא ימין

    for run in p.runs:
        rPr = run._element.get_or_add_rPr()
        if run.bold:
            _insert(rPr, 'w:bCs', RPR_AFTER_BCS)
        if run.font.size:
            _insert(rPr, 'w:szCs', RPR_AFTER_SZCS, int(round(run.font.size.pt * 2)))
        _insert(rPr, 'w:rtl', RPR_AFTER_RTL, 1)


def create_full_word_resume(json_filepath, output_filepath):
    with open(json_filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # זיהוי שפה לפי התוכן (ולא לפי שם הקובץ)
    is_hebrew = any('\u0590' <= ch <= '\u05FF' for ch in data['personal_info']['name'])

    t_summary = 'תמצית' if is_hebrew else 'Summary'
    t_edu = 'השכלה' if is_hebrew else 'Education'
    t_skills = 'מיומנויות וכלים' if is_hebrew else 'Skills'
    t_proj = 'פרויקטים' if is_hebrew else 'Projects'
    t_ml = 'פרויקטים בניתוח נתונים ולמידת מכונה:' if is_hebrew else 'Data Analysis & Machine Learning Projects:'
    t_exp = 'ניסיון תעסוקתי' if is_hebrew else 'Experience'
    t_mil = 'שירות צבאי' if is_hebrew else 'Military Service'
    t_vol = 'מעורבות חברתית והתנדבות' if is_hebrew else 'Social Involvement and Volunteering'
    t_lang = 'שפות' if is_hebrew else 'Languages'

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    style = doc.styles['Normal']
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0

    if 'List Bullet' in doc.styles:
        bullet_style = doc.styles['List Bullet']
        bullet_style.paragraph_format.space_after = Pt(0)
        bullet_style.paragraph_format.space_before = Pt(0)
        bullet_style.paragraph_format.left_indent = Inches(0.25)

    # --- כותרת ופרטי קשר ---
    title_p = doc.add_paragraph()
    run = add_tnr_run(title_p, data['personal_info']['name'], bold=True, size=16)
    run.font.color.rgb = RGBColor(46, 116, 181)
    title_p.paragraph_format.space_after = Pt(2)

    contact_p = doc.add_paragraph()
    contact_p.paragraph_format.space_after = Pt(4)
    add_tnr_run(contact_p, f"{data['personal_info'].get('phone', '')} | {data['personal_info'].get('email', '')}  ")
    add_hyperlink(contact_p, "LinkedIn", data['personal_info'].get('linkedin', ''))
    add_tnr_run(contact_p, " | ")
    add_hyperlink(contact_p, "GitHub", data['personal_info'].get('github', ''))

    # --- תמצית (הנקודה בסוף נשארת) ---
    if 'professional_summary' in data:
        add_tnr_heading(doc, t_summary)
        p = doc.add_paragraph()
        add_tnr_run(p, data['professional_summary'])

    # --- השכלה ---
    if 'education' in data:
        add_tnr_heading(doc, t_edu)
        for edu in data['education']:
            p = doc.add_paragraph()
            add_tnr_run(p, f"{edu['institution']} | {edu['degree']} ", bold=True)
            add_tnr_run(p, f"{edu['years']}")
            if 'notes' in edu:
                for note in edu['notes'].split('. '):
                    if note.strip():
                        add_bullet_point(doc, note, is_hebrew)

    # --- מיומנויות ---
    if 'skills' in data:
        add_tnr_heading(doc, t_skills)
        for category, skills_list in data['skills'].items():
            if category == 'languages' or category == 'software_development_and_ai':
                continue
            clean_category = category.replace('_', ' ').title()
            skills_str = ", ".join(skills_list)
            add_bullet_point(doc, f"{clean_category}: {skills_str}", is_hebrew)

    # --- פרויקטים ---
    if 'projects' in data:
        add_tnr_heading(doc, t_proj)
        ml_header_done = False
        for proj in data['projects']:
            name = proj['name']
            desc = proj['description']
            is_ml = name.startswith(ML_PREFIX) or name[:1].isdigit()

            if is_ml:
                name = name.replace(ML_PREFIX, "")
                if not ml_header_done:
                    p_head = doc.add_paragraph()
                    p_head.paragraph_format.keep_with_next = True
                    add_tnr_run(p_head, t_ml, bold=True)
                    ml_header_done = True
                add_numbered_item(doc, name, desc)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                add_tnr_run(p, name.rstrip(':') + ":", bold=True)
                for part in desc.split('. '):
                    if part.strip():
                        add_bullet_point(doc, part, is_hebrew)

    # --- ניסיון ושירות צבאי ---
    if 'experience' in data:
        add_tnr_heading(doc, t_exp)
        for exp in data['experience']:
            if "Military Service |" in exp['company']:
                add_tnr_heading(doc, t_mil, space_before=4)
                company_name = exp['company'].replace("Military Service | ", "")
                p = doc.add_paragraph()
                add_tnr_run(p, f"{company_name} | {exp['role']} ", bold=True)
                add_tnr_run(p, f"{exp['years']}")
            else:
                p = doc.add_paragraph()
                add_tnr_run(p, f"{exp['company']} | {exp['role']} ", bold=True)
                add_tnr_run(p, f"{exp['years']}")

            for bullet in exp['bullets']:
                add_bullet_point(doc, bullet, is_hebrew)

    # --- התנדבות ---
    if 'volunteering' in data:
        add_tnr_heading(doc, t_vol)
        for vol in data['volunteering']:
            p = doc.add_paragraph()
            add_tnr_run(p, f"{vol['organization']} | {vol['role']} ", bold=True)
            add_tnr_run(p, f"{vol['years']}")
            for bullet in vol['bullets']:
                add_bullet_point(doc, bullet, is_hebrew)

    # --- שפות ---
    if 'skills' in data and 'languages' in data['skills']:
        add_tnr_heading(doc, t_lang)
        for lang in data['skills']['languages']:
            p = doc.add_paragraph()
            add_tnr_run(p, lang)

    # --- כיווניות ויישור לעברית ---
    if is_hebrew:
        for p in doc.paragraphs:
            # שורת הקשר מזוהה לפי האובייקט עצמו (ולא לפי הטקסט!)
            # אחרת שורת "Git/GitHub" במיומנויות נחשבה בטעות כשורת קשר
            if p._p is contact_p._p:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                continue
            make_rtl(p)

    doc.save(output_filepath)
    print(f" נוצר מסמך : {output_filepath}")


if __name__ == "__main__":
    create_full_word_resume('master_resume.json', 'My_Full_Resume.docx')
    convert("My_Full_Resume.docx", "My_Full_Resume.pdf")
    print(" נוצר גם קובץ PDF ")